"""
Business reports: profit & loss summary and CSV / Excel exports.

Definitions (kept simple and printed in the UI):
  revenue        = sum of sale line totals (before tax), voided sales excluded
  cost of goods  = unit cost snapshot taken at the time of each sale
  gross profit   = revenue - cost of goods
  net profit     = gross profit - expenses recorded in the period
"""
import csv
import io
from datetime import timedelta
from decimal import Decimal

from django.db.models import Count, DecimalField, ExpressionWrapper, F, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.utils import timezone

from ..models import CustomerDebt, Expense, Product, Sale
from .analytics import cogs_expression, live_transactions, period_totals
from .formatting import local_now, store_tz

ZERO = Decimal('0.00')


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

def build_summary(store, date_from, date_to, start, end):
    totals = period_totals(store, start, end)

    expenses = Expense.objects.filter(store=store, expense_date__gte=date_from, expense_date__lte=date_to)
    expenses_total = expenses.aggregate(t=Sum('amount'))['t'] or ZERO
    category_labels = dict(Expense.Category.choices)
    expenses_by_category = [
        {
            'category': row['category'],
            'label': category_labels.get(row['category'], row['category']),
            'total': row['total'],
        }
        for row in expenses.values('category').annotate(total=Sum('amount')).order_by('-total')
    ]

    sales = Sale.objects.filter(store=store, sold_at__gte=start, sold_at__lt=end)
    completed = sales.filter(status=Sale.Status.COMPLETED)
    method_labels = dict(Sale.PaymentMethod.choices)
    payment_breakdown = [
        {
            'method': row['payment_method'],
            'label': method_labels.get(row['payment_method'], row['payment_method']),
            'total': row['total'] or ZERO,
            'count': row['count'],
        }
        for row in completed.values('payment_method').annotate(total=Sum('total_amount'), count=Count('id')).order_by('-total')
    ]
    tax_collected = completed.aggregate(t=Sum('tax_amount'))['t'] or ZERO
    voided = sales.filter(status=Sale.Status.VOIDED)
    voided_agg = voided.aggregate(t=Sum('total_amount'), c=Count('id'))

    top_products = []
    top_rows = (
        live_transactions(store)
        .filter(transaction_date__gte=start, transaction_date__lt=end, product__isnull=False)
        .values('product', 'product__name', 'product__sku')
        .annotate(units=Sum('quantity_sold'), revenue=Sum('total_amount'), cogs=Sum(cogs_expression()))
        .order_by('-revenue')[:10]
    )
    for row in top_rows:
        revenue = row['revenue'] or ZERO
        cogs = row['cogs'] or ZERO
        top_products.append({
            'product_id': row['product'],
            'product_name': row['product__name'],
            'sku': row['product__sku'],
            'units': row['units'] or 0,
            'revenue': revenue,
            'profit': revenue - cogs,
        })

    daily_rows = (
        live_transactions(store)
        .filter(transaction_date__gte=start, transaction_date__lt=end)
        .annotate(day=TruncDate('transaction_date', tzinfo=store_tz(store)))
        .values('day')
        .annotate(revenue=Sum('total_amount'), cogs=Sum(cogs_expression()))
    )
    by_day = {row['day']: row for row in daily_rows}
    daily = []
    cursor = date_from
    while cursor <= date_to:
        row = by_day.get(cursor)
        revenue = (row['revenue'] if row else None) or ZERO
        cogs = (row['cogs'] if row else None) or ZERO
        daily.append({'date': cursor.isoformat(), 'revenue': revenue, 'profit': revenue - cogs})
        cursor += timedelta(days=1)

    open_debts = list(CustomerDebt.objects.filter(store=store, is_settled=False))
    inventory_value = Product.objects.filter(store=store, is_active=True).aggregate(
        v=Sum(ExpressionWrapper(
            F('cost_price') * F('current_stock'),
            output_field=DecimalField(max_digits=18, decimal_places=2),
        ))
    )['v'] or ZERO

    orders = totals['orders']
    return {
        'period': {'date_from': date_from.isoformat(), 'date_to': date_to.isoformat()},
        'currency': store.currency,
        'revenue': totals['revenue'],
        'cost_of_goods': totals['cogs'],
        'gross_profit': totals['gross_profit'],
        'expenses_total': expenses_total,
        'net_profit': totals['gross_profit'] - expenses_total,
        'orders': orders,
        'units_sold': totals['units'],
        'average_order_value': (totals['revenue'] / orders).quantize(Decimal('0.01')) if orders else ZERO,
        'tax_collected': tax_collected,
        'voided_count': voided_agg['c'] or 0,
        'voided_total': voided_agg['t'] or ZERO,
        'payment_breakdown': payment_breakdown,
        'expenses_by_category': expenses_by_category,
        'top_products': top_products,
        'daily': daily,
        'debts_outstanding': sum((d.outstanding_amount for d in open_debts), ZERO),
        'debts_overdue': sum((d.outstanding_amount for d in open_debts if d.days_overdue > 0), ZERO),
        'inventory_value': inventory_value,
    }


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def _local_stamp(store, dt):
    if not dt:
        return ''
    return timezone.localtime(dt, store_tz(store)).strftime('%Y-%m-%d %H:%M')


def export_sales(store, start, end, **_):
    sales = (
        Sale.objects.filter(store=store, sold_at__gte=start, sold_at__lt=end)
        .select_related('customer', 'cashier').order_by('sold_at')
    )
    headers = ['Receipt', 'Date', 'Customer', 'Cashier', 'Payment method', 'Status',
               'Subtotal', 'Tax', 'Total', 'Amount paid']
    rows = [
        [
            s.receipt_number, _local_stamp(store, s.sold_at),
            s.customer.name if s.customer else '',
            (s.cashier.full_name or s.cashier.username) if s.cashier else '',
            s.get_payment_method_display(), s.get_status_display(),
            s.subtotal, s.tax_amount, s.total_amount, s.amount_paid,
        ]
        for s in sales
    ]
    return 'sales', headers, rows


def export_sale_items(store, start, end, **_):
    lines = (
        live_transactions(store)
        .filter(transaction_date__gte=start, transaction_date__lt=end)
        .select_related('product', 'sale', 'customer').order_by('transaction_date')
    )
    headers = ['Date', 'Receipt', 'Product', 'SKU', 'Quantity', 'Unit price',
               'Line total', 'Unit cost', 'Profit', 'Customer']
    rows = []
    for t in lines:
        cost = (t.unit_cost or ZERO) * t.quantity_sold
        rows.append([
            _local_stamp(store, t.transaction_date),
            t.sale.receipt_number if t.sale else '',
            t.product.name if t.product else '',
            t.product.sku if t.product else '',
            t.quantity_sold, t.unit_price, t.total_amount, t.unit_cost or ZERO,
            t.total_amount - cost,
            t.customer.name if t.customer else '',
        ])
    return 'sale-items', headers, rows


def export_inventory(store, **_):
    products = Product.objects.filter(store=store, is_active=True).order_by('name')
    headers = ['SKU', 'Product', 'Category', 'Unit', 'In stock', 'Reorder point',
               'Cost price', 'Selling price', 'Stock value (at cost)', 'Margin %']
    rows = [
        [
            p.sku, p.name, p.category, p.unit_of_measure, p.current_stock, p.reorder_point,
            p.cost_price, p.selling_price, p.cost_price * p.current_stock, round(p.margin_percent, 1),
        ]
        for p in products
    ]
    return 'inventory', headers, rows


def export_debtors(store, **_):
    debts = CustomerDebt.objects.filter(store=store, is_settled=False).select_related('customer').order_by('due_date')
    headers = ['Customer', 'Phone', 'Original amount', 'Paid', 'Outstanding', 'Due date', 'Days overdue']
    rows = [
        [
            d.customer.name, d.customer.phone_number, d.amount, d.amount_paid,
            d.outstanding_amount, d.due_date.isoformat(), d.days_overdue,
        ]
        for d in debts
    ]
    return 'debtors', headers, rows


def export_expenses(store, date_from=None, date_to=None, **_):
    expenses = (
        Expense.objects.filter(store=store, expense_date__gte=date_from, expense_date__lte=date_to)
        .select_related('recorded_by').order_by('expense_date')
    )
    headers = ['Date', 'Category', 'Description', 'Amount', 'Payment method', 'Recorded by']
    rows = [
        [
            e.expense_date.isoformat(), e.get_category_display(), e.description, e.amount,
            e.get_payment_method_display(),
            (e.recorded_by.full_name or e.recorded_by.username) if e.recorded_by else '',
        ]
        for e in expenses
    ]
    return 'expenses', headers, rows


EXPORTERS = {
    'sales': export_sales,
    'sale-items': export_sale_items,
    'inventory': export_inventory,
    'debtors': export_debtors,
    'expenses': export_expenses,
}


def _safe_cell(value):
    """Stop spreadsheet formula injection: a text cell starting with = + - @ is neutralised."""
    if isinstance(value, str) and value and value[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + value
    return value


def render_export(store, kind, fmt, date_from, date_to, start, end):
    """Build an HttpResponse containing the requested export as CSV or XLSX."""
    name, headers, rows = EXPORTERS[kind](
        store, start=start, end=end, date_from=date_from, date_to=date_to,
    )
    stamp = local_now(store).strftime('%Y%m%d')
    filename = f"{name}-{stamp}.{fmt}"
    clean_rows = [[_safe_cell(v) for v in row] for row in rows]

    if fmt == 'csv':
        buffer = io.StringIO()
        buffer.write('\ufeff')  # BOM so Excel opens UTF-8 (e.g. the naira sign) correctly
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(clean_rows)
        response = HttpResponse(buffer.getvalue(), content_type='text/csv; charset=utf-8')
    else:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = name[:31]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in clean_rows:
            sheet.append(row)
        for idx, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(r[idx - 1])) for r in clean_rows[:200]])
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 50)
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
