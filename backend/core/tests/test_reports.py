"""Profit & loss summary and CSV / Excel exports."""
import io
from decimal import Decimal

import openpyxl

from core.models import Expense
from core.services.formatting import local_now
from core.services.reports import _safe_cell
from .helpers import ShopTestCase


class ReportTests(ShopTestCase):

    def setUp(self):
        super().setUp()
        # 4 x A (price 100, cost 50) = revenue 400, cost 200; 1 x B (200 / 120) = revenue 200, cost 120
        self.sell([(self.product_a, 4), (self.product_b, 1)], customer=str(self.customer.id))
        Expense.objects.create(
            store=self.store, category='RENT', description='Rent', amount=Decimal('100.00'),
            expense_date=local_now(self.store).date(),
        )

    def test_summary_numbers(self):
        response = self.api('get', 'reports/summary/')
        self.assertEqual(response.status_code, 200, response.data)
        data = response.data
        self.assertEqual(Decimal(data['revenue']), Decimal('600.00'))
        self.assertEqual(Decimal(data['cost_of_goods']), Decimal('320.00'))
        self.assertEqual(Decimal(data['gross_profit']), Decimal('280.00'))
        self.assertEqual(Decimal(data['expenses_total']), Decimal('100.00'))
        self.assertEqual(Decimal(data['net_profit']), Decimal('180.00'))
        self.assertEqual(Decimal(data['tax_collected']), Decimal('60.00'))     # 10% of 600
        self.assertEqual(data['orders'], 1)
        self.assertEqual(data['units_sold'], 5)
        self.assertEqual(data['top_products'][0]['product_name'], 'Product A')   # 400 > 200
        self.assertEqual(data['payment_breakdown'][0]['method'], 'CASH')
        self.assertEqual(data['expenses_by_category'][0]['category'], 'RENT')
        self.assertEqual(len(data['daily']), 30)

    def test_voided_sales_leave_the_summary(self):
        sale_id = self.results(self.api('get', 'sales/'))[0]['id']
        self.api('post', f'sales/{sale_id}/void/')
        data = self.api('get', 'reports/summary/').data
        self.assertEqual(Decimal(data['revenue']), Decimal('0.00'))
        self.assertEqual(data['voided_count'], 1)

    def test_period_validation(self):
        self.assertEqual(self.api('get', 'reports/summary/', {'date_from': 'abc'}).status_code, 400)
        self.assertEqual(
            self.api('get', 'reports/summary/', {'date_from': '2026-03-01', 'date_to': '2026-01-01'}).status_code, 400
        )
        self.assertEqual(
            self.api('get', 'reports/summary/', {'date_from': '2020-01-01', 'date_to': '2026-01-01'}).status_code, 400
        )

    def test_roles(self):
        for user, expected in (
            (self.accountant, 200), (self.manager, 200), (self.cashier, 403), (self.viewer, 403),
        ):
            self.client.force_authenticate(user=user)
            self.assertEqual(self.api('get', 'reports/summary/').status_code, expected, user.username)

    def test_csv_export(self):
        response = self.api('get', 'reports/export/sales/', {'fmt': 'csv'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        self.assertIn('attachment; filename="sales-', response['Content-Disposition'])
        text = response.content.decode('utf-8-sig')
        self.assertTrue(text.startswith('Receipt,Date,Customer'))
        self.assertIn('RCP-000001', text)

    def test_xlsx_export(self):
        response = self.api('get', 'reports/export/inventory/', {'fmt': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet['A1'].value, 'SKU')
        skus = {sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)}
        self.assertEqual(skus, {'A-1', 'B-1'})

    def test_every_export_kind_works_in_both_formats(self):
        for kind in ('sales', 'sale-items', 'inventory', 'debtors', 'expenses'):
            for fmt in ('csv', 'xlsx'):
                response = self.api('get', f'reports/export/{kind}/', {'fmt': fmt})
                self.assertEqual(response.status_code, 200, f'{kind}.{fmt}')

    def test_unknown_report_and_bad_format(self):
        self.assertEqual(self.api('get', 'reports/export/secrets/').status_code, 404)
        self.assertEqual(self.api('get', 'reports/export/sales/', {'fmt': 'pdf'}).status_code, 400)

    def test_formula_injection_is_neutralised(self):
        self.assertEqual(_safe_cell('=HYPERLINK("http://evil")'), "'=HYPERLINK(\"http://evil\")")
        self.assertEqual(_safe_cell('+1'), "'+1")
        self.assertEqual(_safe_cell('Normal'), 'Normal')
        self.assertEqual(_safe_cell(-5), -5)
